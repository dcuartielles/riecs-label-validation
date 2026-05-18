import io
from collections import defaultdict
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from fastapi import APIRouter, Request
from fastapi.responses import StreamingResponse, RedirectResponse
from sqlalchemy import select, func
from sqlalchemy.orm import selectinload

from app.auth import get_current_user
from app.database import SessionLocal
from app.models import (
    AddedLabel, AddedLabelDecision, Group, GroupAssignment,
    MandatoryClassification, Session as ReviewSession,
    StoryRejection, StoryRelevance, TaxonomyLabel, User, UserStory,
)

router = APIRouter()

from app.config import find_latest_dataset
DATASET_PATH = find_latest_dataset()
LABELBOOK_PATH = Path("labelbook/Labelbook 2026-05-05_used.xlsx")

FILL_YELLOW   = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
FILL_GREEN    = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FILL_RED      = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
FILL_LT_GREEN = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
FILL_LT_RED   = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
FILL_HDR      = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
FILL_HDR_NEW  = PatternFill(start_color="375623", end_color="375623", fill_type="solid")
FILL_STAT_HDR = PatternFill(start_color="2C324C", end_color="2C324C", fill_type="solid")
FILL_ORANGE   = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
FILL_AMBER    = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
FILL_BLUE_HDR = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
FILL_RED_HDR  = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")

FONT_WHITE_BOLD  = Font(bold=True, color="FFFFFF")
FILL_DARK_RED    = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
FILL_NEW_LABEL   = PatternFill(start_color="FFD280", end_color="FFD280", fill_type="solid")

STORY_COLS = ["Story ID", "Workshop", "Submitted by", "Stakeholder Group",
              "User type", "Task", "Goal", "Additional Notes"]


def _story_row(story: UserStory) -> list:
    return [
        story.story_id,
        story.workshop or "",
        story.submitted_by or "",
        story.stakeholder_group or "",
        story.user_type or "",
        story.task or "",
        story.goal or "",
        story.additional_notes or "",
    ]


def _write_header_row(ws, cols: list[str], fill=None):
    for c, val in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c, value=val)
        cell.font = FONT_WHITE_BOLD if fill else Font(bold=True)
        if fill:
            cell.fill = fill


def _autofit(ws, max_width=60):
    for col in ws.columns:
        w = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(w + 2, max_width)


def build_partner_sheet(wb_out, group_name: str, stories: list[UserStory],
                         added_map: dict, mandatory_map: dict,
                         rejection_map: dict, relevance_map: dict,
                         max_labels: int):
    """One sheet per partner with story data + labelling results."""
    safe_name = group_name[:31]  # Excel sheet name limit
    ws = wb_out.create_sheet(title=safe_name)

    extra_cols = ["Target user", "Concepts", "Rejected", "Rejection reason",
                  "Relevance", "Relevance reason"]
    label_cols = [f"Label {i+1}" for i in range(max_labels)]
    all_cols = STORY_COLS + label_cols + extra_cols

    _write_header_row(ws, all_cols, FILL_STAT_HDR)

    # Colour the label header cells differently
    for i in range(max_labels):
        cell = ws.cell(row=1, column=len(STORY_COLS) + i + 1)
        cell.fill = FILL_HDR

    for row_idx, story in enumerate(stories, start=2):
        sid = story.id
        base = _story_row(story)

        labels = added_map.get(sid, [])
        label_vals = [
            (lbl["label"] + (" › " + lbl["sublabel"] if lbl.get("sublabel") else "")
             + (" (" + lbl["note"] + ")" if lbl.get("note") else ""))
            for lbl in labels
        ]
        label_vals += [""] * (max_labels - len(label_vals))

        mc = mandatory_map.get(sid)
        rej = rejection_map.get(sid)
        rel = relevance_map.get(sid)

        extra = [
            mc.target_user if mc and mc.target_user else "",
            "; ".join(mc.concepts) if mc and mc.concepts else "",
            "Yes" if rej and rej.rejected else "No",
            rej.reason if rej and rej.reason else "",
            rel.score if rel else "Normal",
            rel.reason if rel and rel.reason else "",
        ]

        full_row = base + label_vals + extra
        for c, val in enumerate(full_row, 1):
            cell = ws.cell(row=row_idx, column=c, value=val)
            col_idx = c
            # Colour label cells green
            if len(STORY_COLS) < col_idx <= len(STORY_COLS) + max_labels and val:
                cell.fill = FILL_GREEN
            # Colour rejected rows light red
            if rej and rej.rejected:
                if col_idx <= len(STORY_COLS):
                    cell.fill = FILL_LT_RED

    _autofit(ws)
    return ws


def build_summary_sheet(wb_out, groups: list, stories: list[UserStory],
                         per_group_added: dict, per_group_mandatory: dict,
                         per_group_rejection: dict, per_group_relevance: dict):
    """
    Summary sheet: each row = one story, columns = per-group label count,
    rejection, relevance. Rows with 2+ groups sorted by partner count descending.
    """
    ws = wb_out.create_sheet(title="Summary")

    group_names = [g.name for g in groups]
    # Per-group columns: label count, rejected, relevance
    extra_headers = []
    for gn in group_names:
        short = gn.split("–")[0].strip() if "–" in gn else gn[:20]
        extra_headers += [f"{short} labels", f"{short} rejected", f"{short} relevance"]

    all_cols = STORY_COLS + extra_headers + ["Groups reviewed", "Has conflict"]
    _write_header_row(ws, all_cols, FILL_STAT_HDR)

    # Build rows
    row_data = []
    for story in stories:
        sid = story.id
        groups_reviewed = 0
        extra_vals = []
        group_labels: dict[str, list] = {}

        for group in groups:
            gid = group.id
            added = per_group_added.get(gid, {}).get(sid, [])
            rej = per_group_rejection.get(gid, {}).get(sid)
            rel = per_group_relevance.get(gid, {}).get(sid)

            n_labels = len(added)
            if n_labels > 0 or rej is not None or rel is not None:
                groups_reviewed += 1

            group_labels[group.name] = [lbl["label"] + ("/" + lbl["sublabel"] if lbl.get("sublabel") else "")
                                         for lbl in added]

            extra_vals.append(n_labels if n_labels else "")
            extra_vals.append("Yes" if rej and rej.rejected else ("" if rej is None else "No"))
            extra_vals.append(rel.score if rel else "")

        # Detect conflict: same story, different labels across groups
        all_labels_sets = [set(lbls) for lbls in group_labels.values() if lbls]
        has_conflict = False
        if len(all_labels_sets) >= 2:
            union = set.union(*all_labels_sets)
            intersect = set.intersection(*all_labels_sets)
            has_conflict = len(union) > len(intersect)

        row_data.append({
            "story": story,
            "groups_reviewed": groups_reviewed,
            "has_conflict": has_conflict,
            "extra_vals": extra_vals,
        })

    # Sort: multi-group stories first, then by groups_reviewed desc
    row_data.sort(key=lambda r: (-r["groups_reviewed"], r["story"].story_id))

    for row_idx, rd in enumerate(row_data, start=2):
        story = rd["story"]
        full_row = _story_row(story) + rd["extra_vals"] + [rd["groups_reviewed"], "Yes" if rd["has_conflict"] else ""]
        for c, val in enumerate(full_row, 1):
            cell = ws.cell(row=row_idx, column=c, value=val)
            # Conflict rows: light orange
            if rd["has_conflict"] and rd["groups_reviewed"] >= 2 and c <= len(STORY_COLS):
                cell.fill = FILL_ORANGE
            # Multi-group non-conflict: light green
            elif not rd["has_conflict"] and rd["groups_reviewed"] >= 2 and c <= len(STORY_COLS):
                cell.fill = FILL_LT_GREEN

    _autofit(ws)
    return ws


def build_rejection_relevance_sheet(wb_out, groups, all_stories,
                                     per_group_rejection, per_group_relevance,
                                     gid_name: dict, uid_name: dict):
    """Stories flagged as rejected or elevated relevance, one row per partner per story."""
    ws = wb_out.create_sheet(title="Rejections & Relevance")
    cols = ["Story ID", "Workshop", "User type", "Task",
            "Partner", "Actor", "Rejected", "Rejection reason", "Relevance", "Relevance reason"]
    _write_header_row(ws, cols, FILL_RED_HDR)

    sid_story = {s.id: s for s in all_stories}
    rows = []

    for group in groups:
        gid = group.id
        all_sids = set(per_group_rejection.get(gid, {}).keys()) | set(per_group_relevance.get(gid, {}).keys())
        for sid in all_sids:
            rej = per_group_rejection.get(gid, {}).get(sid)
            rel = per_group_relevance.get(gid, {}).get(sid)
            is_rejected = bool(rej and rej.rejected)
            score = rel.score if rel else "Normal"
            if not is_rejected and score == "Normal":
                continue
            story = sid_story.get(sid)
            if not story:
                continue
            actor_uid = (rej.user_id if rej else None) or (rel.user_id if rel else None)
            rows.append({
                "story": story,
                "partner": gid_name.get(gid, "?"),
                "actor": uid_name.get(actor_uid, "?") if actor_uid else "",
                "rejected": "Yes" if is_rejected else "No",
                "rej_reason": rej.reason if rej and rej.reason else "",
                "score": score,
                "rel_reason": rel.reason if rel and rel.reason else "",
                "sort_key": (0 if is_rejected else 1,
                             {"VeryHigh": 0, "High": 1}.get(score, 2),
                             story.story_id or ""),
            })

    rows.sort(key=lambda r: r["sort_key"])

    for row_idx, rd in enumerate(rows, start=2):
        s = rd["story"]
        vals = [s.story_id, s.workshop or "", s.user_type or "", (s.task or "")[:120],
                rd["partner"], rd["actor"], rd["rejected"], rd["rej_reason"],
                rd["score"], rd["rel_reason"]]
        for c, val in enumerate(vals, 1):
            cell = ws.cell(row=row_idx, column=c, value=val)
            if rd["rejected"] == "Yes":
                cell.fill = FILL_DARK_RED
                cell.font = FONT_WHITE_BOLD
            elif rd["score"] == "VeryHigh":
                cell.fill = FILL_ORANGE
            elif rd["score"] == "High":
                cell.fill = FILL_AMBER

    _autofit(ws)
    return ws


def build_stats_conflicts_sheet(wb_out, groups, all_stories,
                                  per_group_added, per_group_mandatory,
                                  per_group_rejection, per_group_relevance,
                                  assigned_counts: dict):
    """Per-partner statistics table + conflict list for stories reviewed by 2+ partners."""
    ws = wb_out.create_sheet(title="Statistics & Conflicts")

    # ── Section 1: per-partner statistics ──────────────────────────────────
    title_cell = ws.cell(row=1, column=1, value="Per-partner statistics")
    title_cell.font = Font(bold=True, size=13, color="FFFFFF")
    title_cell.fill = FILL_BLUE_HDR

    stat_cols = ["Partner", "Stories assigned", "Stories labelled",
                 "Labels added", "Mandatory filled", "Rejected", "High/VH relevance"]
    for c, col in enumerate(stat_cols, 1):
        cell = ws.cell(row=2, column=c, value=col)
        cell.font = FONT_WHITE_BOLD
        cell.fill = FILL_STAT_HDR

    for r, group in enumerate(groups, start=3):
        gid = group.id
        added    = per_group_added.get(gid, {})
        mandatory = per_group_mandatory.get(gid, {})
        rejection = per_group_rejection.get(gid, {})
        relevance = per_group_relevance.get(gid, {})

        vals = [
            group.name,
            assigned_counts.get(gid, 0),
            sum(1 for lbls in added.values() if lbls),
            sum(len(lbls) for lbls in added.values()),
            len(mandatory),
            sum(1 for rj in rejection.values() if rj.rejected),
            sum(1 for rv in relevance.values() if rv.score in ("High", "VeryHigh")),
        ]
        for c, val in enumerate(vals, 1):
            ws.cell(row=r, column=c, value=val)

    gap = len(groups) + 4

    # ── Section 2: conflicts ────────────────────────────────────────────────
    title_cell2 = ws.cell(row=gap, column=1, value="Conflicts (stories reviewed by 2+ partners)")
    title_cell2.font = Font(bold=True, size=13, color="FFFFFF")
    title_cell2.fill = FILL_RED_HDR

    conflict_cols = ["Story ID", "Workshop", "User type", "Task",
                     "Partners reviewed", "Conflict on"] + [g.name for g in groups]
    for c, col in enumerate(conflict_cols, 1):
        cell = ws.cell(row=gap + 1, column=c, value=col)
        cell.font = FONT_WHITE_BOLD
        cell.fill = FILL_STAT_HDR if c <= 6 else FILL_HDR

    conflict_rows = []
    for story in all_stories:
        sid = story.id
        groups_with_data = [
            g for g in groups
            if per_group_added.get(g.id, {}).get(sid)
            or sid in per_group_mandatory.get(g.id, {})
        ]
        if len(groups_with_data) < 2:
            continue

        tu_vals = {
            per_group_mandatory.get(g.id, {}).get(sid).target_user
            for g in groups_with_data
            if per_group_mandatory.get(g.id, {}).get(sid) and
               per_group_mandatory[g.id][sid].target_user
        }
        lbl_sets = [
            frozenset(
                (lbl["sublabel"] or lbl["label"])
                for lbl in per_group_added.get(g.id, {}).get(sid, [])
            )
            for g in groups_with_data
        ]
        tu_conflict  = len(tu_vals) > 1
        lbl_union    = set.union(*[set(s) for s in lbl_sets]) if lbl_sets else set()
        lbl_intersect = set.intersection(*[set(s) for s in lbl_sets]) if lbl_sets else set()
        lbl_conflict = len(lbl_union) > len(lbl_intersect)

        if not (tu_conflict or lbl_conflict):
            continue

        conflict_rows.append({
            "story": story,
            "n_partners": len(groups_with_data),
            "conflict_on": ", ".join(
                (["Target user"] if tu_conflict else []) +
                (["Labels"] if lbl_conflict else [])
            ),
        })

    conflict_rows.sort(key=lambda r: -r["n_partners"])

    for row_idx, rd in enumerate(conflict_rows, start=gap + 2):
        s = rd["story"]
        base = [s.story_id, s.workshop or "", s.user_type or "",
                (s.task or "")[:100], rd["n_partners"], rd["conflict_on"]]
        for c, val in enumerate(base, 1):
            cell = ws.cell(row=row_idx, column=c, value=val)
            cell.fill = FILL_ORANGE

        for g_idx, group in enumerate(groups):
            mc   = per_group_mandatory.get(group.id, {}).get(s.id)
            lbls = per_group_added.get(group.id, {}).get(s.id, [])
            if mc or lbls:
                tu_str  = mc.target_user if mc else ""
                lbl_str = ", ".join(
                    (lbl["sublabel"] or lbl["label"]) for lbl in lbls
                )
                cell_val = ("TU: " + tu_str if tu_str else "") + \
                           ("\nLabels: " + lbl_str if lbl_str else "")
            else:
                cell_val = ""
            cell = ws.cell(row=row_idx, column=7 + g_idx, value=cell_val)
            if cell_val:
                cell.alignment = Alignment(wrap_text=True)

    _autofit(ws)
    return ws


def build_label_authorship_sheet(wb_out, authorship_rows, uid_name, uid_group, gid_name):
    """One row per added label: who added what to which story."""
    ws = wb_out.create_sheet(title="Label Authorship")
    cols = ["Story ID", "Workshop", "User Type", "Task",
            "Group", "Added By", "Label", "Sublabel", "Note", "New?"]
    _write_header_row(ws, cols, FILL_STAT_HDR)

    for row_idx, al in enumerate(authorship_rows, start=2):
        s = al.story
        tl = al.taxonomy_label
        label_str   = tl.label    if tl else (al.free_text or "")
        sublabel_str = tl.sublabel if tl else ""
        is_new = "(*)" if (tl and tl.is_user_created) else ""
        group_name = gid_name.get(uid_group.get(al.user_id), "?")
        vals = [
            s.story_id if s else "",
            s.workshop or "" if s else "",
            s.user_type or "" if s else "",
            (s.task or "")[:120] if s else "",
            group_name,
            uid_name.get(al.user_id, "?"),
            label_str,
            sublabel_str,
            al.note or "",
            is_new,
        ]
        for c, val in enumerate(vals, 1):
            cell = ws.cell(row=row_idx, column=c, value=val)
            if is_new:
                cell.fill = FILL_LT_GREEN

    _autofit(ws)
    return ws


def build_label_decisions_sheet(wb_out, decision_rows, uid_name, uid_group, gid_name):
    """Who confirmed or rejected each peer label."""
    ws = wb_out.create_sheet(title="Label Decisions")
    cols = ["Story ID", "Group", "Label", "Sublabel",
            "Added By", "Decision", "Decided By"]
    _write_header_row(ws, cols, FILL_STAT_HDR)

    for row_idx, dec in enumerate(decision_rows, start=2):
        al = dec.added_label
        tl = al.taxonomy_label if al else None
        s  = al.story if al else None
        label_str    = tl.label    if tl else (al.free_text if al else "")
        sublabel_str = tl.sublabel if tl else ""
        adder_group  = gid_name.get(uid_group.get(al.user_id if al else None), "?")
        vals = [
            s.story_id if s else "",
            adder_group,
            label_str,
            sublabel_str,
            uid_name.get(al.user_id if al else None, "?"),
            dec.decision,
            uid_name.get(dec.user_id, "?"),
        ]
        for c, val in enumerate(vals, 1):
            cell = ws.cell(row=row_idx, column=c, value=val)
            if dec.decision == "reject":
                cell.fill = FILL_DARK_RED
                cell.font = FONT_WHITE_BOLD
            elif dec.decision == "confirm":
                cell.fill = FILL_LT_GREEN

    _autofit(ws)
    return ws


async def generate_workbook(
    db,
    session_id: int | None = None,
    group_ids: list[int] | None = None,
) -> tuple[openpyxl.Workbook, "ReviewSession | None"]:
    """Build the export workbook from an open DB session.

    Callable from the HTTP route, the session-end handler, and the daily
    export script without duplicating any data-loading logic.
    group_ids: if given, restrict to those groups (non-admin user view).
    Returns (workbook, review_session).
    """
    groups = (await db.execute(select(Group))).scalars().all()
    if group_ids is not None:
        groups = [g for g in groups if g.id in group_ids]

    if session_id:
        rev_session = await db.get(ReviewSession, session_id)
    else:
        result = await db.execute(
            select(ReviewSession).order_by(ReviewSession.started_at.desc())
        )
        rev_session = result.scalars().first()

    all_stories = (await db.execute(
        select(UserStory).order_by(UserStory.story_id)
    )).scalars().all()
    story_by_id = {s.id: s for s in all_stories}

    all_users = (await db.execute(select(User))).scalars().all()
    uid_name  = {u.id: u.name     for u in all_users}
    uid_group = {u.id: u.group_id for u in all_users}
    gid_name  = {g.id: g.name     for g in groups}

    per_group_added:     dict[int, dict[int, list]] = {}
    per_group_mandatory: dict[int, dict[int, MandatoryClassification]] = {}
    per_group_rejection: dict[int, dict[int, StoryRejection]] = {}
    per_group_relevance: dict[int, dict[int, StoryRelevance]] = {}

    for group in groups:
        gid = group.id
        added_filter = [User.group_id == gid]
        if rev_session:
            added_filter.append(AddedLabel.session_id == rev_session.id)

        added_result = await db.execute(
            select(AddedLabel)
            .join(User, AddedLabel.user_id == User.id)
            .where(*added_filter)
            .options(selectinload(AddedLabel.taxonomy_label))
        )
        added_labels = added_result.scalars().all()
        added_map: dict[int, list] = {}
        for al in added_labels:
            entry = {
                "label":    al.taxonomy_label.label    if al.taxonomy_label else (al.free_text or ""),
                "sublabel": al.taxonomy_label.sublabel if al.taxonomy_label else "",
                "note":     al.note or "",
            }
            added_map.setdefault(al.story_id, []).append(entry)
        per_group_added[gid] = added_map

        mc_filter = [MandatoryClassification.group_id == gid]
        if rev_session:
            mc_filter.append(MandatoryClassification.session_id == rev_session.id)
        mc_rows = (await db.execute(select(MandatoryClassification).where(*mc_filter))).scalars().all()
        per_group_mandatory[gid] = {mc.story_id: mc for mc in mc_rows}

        rej_filter = [StoryRejection.group_id == gid]
        if rev_session:
            rej_filter.append(StoryRejection.session_id == rev_session.id)
        rej_rows = (await db.execute(select(StoryRejection).where(*rej_filter))).scalars().all()
        per_group_rejection[gid] = {r.story_id: r for r in rej_rows}

        rel_filter = [StoryRelevance.group_id == gid]
        if rev_session:
            rel_filter.append(StoryRelevance.session_id == rev_session.id)
        rel_rows = (await db.execute(select(StoryRelevance).where(*rel_filter))).scalars().all()
        per_group_relevance[gid] = {r.story_id: r for r in rel_rows}

    assigned_counts: dict[int, int] = {}
    for group in groups:
        res = await db.execute(
            select(func.count(GroupAssignment.id))
            .where(GroupAssignment.group_id == group.id)
        )
        assigned_counts[group.id] = res.scalar() or 0

    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)

    for group in groups:
        gid = group.id
        assigned_result = await db.execute(
            select(GroupAssignment.story_id)
            .where(GroupAssignment.group_id == gid)
            .order_by(GroupAssignment.position)
        )
        assigned_ids = [r for r, in assigned_result.all()]
        partner_stories = (
            [story_by_id[sid] for sid in assigned_ids if sid in story_by_id]
            if assigned_ids else all_stories
        )
        max_labels_group = max(
            (len(lbls) for sid, lbls in per_group_added[gid].items()),
            default=1,
        )
        build_partner_sheet(
            wb_out, group.name, partner_stories,
            per_group_added[gid], per_group_mandatory[gid],
            per_group_rejection[gid], per_group_relevance[gid],
            max_labels_group,
        )

    build_summary_sheet(
        wb_out, groups, all_stories,
        per_group_added, per_group_mandatory,
        per_group_rejection, per_group_relevance,
    )
    build_rejection_relevance_sheet(
        wb_out, groups, all_stories,
        per_group_rejection, per_group_relevance,
        gid_name, uid_name,
    )
    build_stats_conflicts_sheet(
        wb_out, groups, all_stories,
        per_group_added, per_group_mandatory,
        per_group_rejection, per_group_relevance,
        assigned_counts,
    )

    # Label authorship
    auth_filter = []
    if rev_session:
        auth_filter.append(AddedLabel.session_id == rev_session.id)
    if group_ids is not None:
        auth_filter.append(User.group_id.in_(group_ids))
    authorship_rows = (await db.execute(
        select(AddedLabel)
        .join(User, AddedLabel.user_id == User.id)
        .options(
            selectinload(AddedLabel.story),
            selectinload(AddedLabel.taxonomy_label),
        )
        .where(*auth_filter)
        .order_by(AddedLabel.story_id, AddedLabel.user_id)
    )).scalars().all()
    build_label_authorship_sheet(wb_out, authorship_rows, uid_name, uid_group, gid_name)

    # Label decisions (peer confirm/reject)
    dec_filter = []
    if rev_session:
        dec_filter.append(AddedLabelDecision.session_id == rev_session.id)
    decision_rows = (await db.execute(
        select(AddedLabelDecision)
        .options(
            selectinload(AddedLabelDecision.added_label).selectinload(AddedLabel.story),
            selectinload(AddedLabelDecision.added_label).selectinload(AddedLabel.taxonomy_label),
        )
        .where(*dec_filter)
        .order_by(AddedLabelDecision.added_label_id)
    )).scalars().all()
    build_label_decisions_sheet(wb_out, decision_rows, uid_name, uid_group, gid_name)

    return wb_out, rev_session


@router.get("/export")
async def export_all(request: Request, session_id: int | None = None):
    user = await get_current_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)

    async with SessionLocal() as db:
        group_ids = None if user.is_admin else [user.group_id]
        wb_out, rev_session = await generate_workbook(db, session_id, group_ids)

    session_suffix = f"session_{rev_session.id}" if rev_session else "all"
    filename = f"label_results_{session_suffix}.xlsx"

    buf = io.BytesIO()
    wb_out.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/export/labelbook")
async def export_labelbook(request: Request):
    user = await get_current_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=302)

    async with SessionLocal() as db:
        # New labels sorted by first appearance (AddedLabel.id = arrival order)
        arrival_rows = (await db.execute(
            select(TaxonomyLabel.id, func.min(AddedLabel.id).label("first_id"))
            .join(AddedLabel, AddedLabel.taxonomy_label_id == TaxonomyLabel.id)
            .where(TaxonomyLabel.is_user_created == True)
            .group_by(TaxonomyLabel.id)
            .order_by(func.min(AddedLabel.id))
        )).all()
        arrival_ids = [tid for tid, _ in arrival_rows]

        tl_by_id = {
            t.id: t for t in (await db.execute(
                select(TaxonomyLabel).where(TaxonomyLabel.is_user_created == True)
            )).scalars().all()
        }
        new_labels = [tl_by_id[tid] for tid in arrival_ids if tid in tl_by_id]

        creators: dict[int, str] = {}
        if arrival_ids:
            for tid, uname in (await db.execute(
                select(AddedLabel.taxonomy_label_id, User.name)
                .join(User, AddedLabel.user_id == User.id)
                .where(AddedLabel.taxonomy_label_id.in_(arrival_ids))
                .order_by(AddedLabel.id)
            )).all():
                creators.setdefault(tid, uname)

    BODY_SIZE = 10
    _body_font = Font(size=BODY_SIZE)

    def _nl(ws, r, c, val=""):
        cell = ws.cell(r, c, val)
        cell.fill = FILL_NEW_LABEL
        cell.font = _body_font
        return cell

    def _copy_row(src_ws, dst_ws, src_r, dst_r, n_cols):
        for c in range(1, n_cols + 1):
            src = src_ws.cell(src_r, c)
            dst = dst_ws.cell(dst_r, c)
            dst.value = src.value
            if src.has_style:
                f = copy(src.font)
                dst.font = Font(name=f.name, size=BODY_SIZE, bold=f.bold,
                                italic=f.italic, color=f.color, underline=f.underline)
                dst.fill      = copy(src.fill)
                dst.border    = copy(src.border)
                dst.alignment = copy(src.alignment)
            else:
                dst.font = Font(size=BODY_SIZE)

    wb = openpyxl.load_workbook(LABELBOOK_PATH)

    # ── Sheet 1: Original Labelbook (untouched) ───────────────────────────
    ws_orig = wb.active
    ws_orig.title = "Original Labelbook"
    n_cols = max(ws_orig.max_column, 7)

    # Parse original structure: group header rows and sublabel rows
    HEADER_ROW = 1
    current_group = None
    group_order: list[str] = []
    group_header_row: dict[str, int] = {}
    group_sublabel_rows: dict[str, list[int]] = defaultdict(list)

    for r in range(HEADER_ROW + 1, ws_orig.max_row + 1):
        v = ws_orig.cell(r, 3).value
        if v:
            current_group = v
            if v not in group_order:
                group_order.append(v)
                group_header_row[v] = r
        elif current_group:
            group_sublabel_rows[current_group].append(r)

    # Group new labels by top-level category, preserving arrival order
    new_by_group: dict[str, list] = defaultdict(list)
    for t in new_labels:
        new_by_group[t.label].append(t)

    # ── Sheet 2: Revised Labelbook (new labels inserted into groups) ───────
    ws_rev = wb.create_sheet("Revised Labelbook", 1)
    _copy_row(ws_orig, ws_rev, HEADER_ROW, HEADER_ROW, n_cols)
    ws_rev.cell(HEADER_ROW, 7).value = "Created by"

    wr = HEADER_ROW + 1
    for group_name in group_order:
        _copy_row(ws_orig, ws_rev, group_header_row[group_name], wr, n_cols)
        wr += 1
        for r in group_sublabel_rows[group_name]:
            _copy_row(ws_orig, ws_rev, r, wr, n_cols)
            wr += 1
        for t in new_by_group.get(group_name, []):
            _nl(ws_rev, wr, 4, t.sublabel)
            _nl(ws_rev, wr, 5, t.source or "")
            _nl(ws_rev, wr, 6, t.description or "")
            _nl(ws_rev, wr, 7, creators.get(t.id, ""))
            wr += 1

    # New labels whose top-level group is not in the original
    for group_name, labels in new_by_group.items():
        if group_name not in group_order:
            _nl(ws_rev, wr, 3, group_name)
            wr += 1
            for t in labels:
                _nl(ws_rev, wr, 4, t.sublabel)
                _nl(ws_rev, wr, 5, t.source or "")
                _nl(ws_rev, wr, 6, t.description or "")
                _nl(ws_rev, wr, 7, creators.get(t.id, ""))
                wr += 1

    _autofit(ws_rev)

    # ── Sheet 3: New Labels only ──────────────────────────────────────────
    ws_new = wb.create_sheet("New Labels")
    _write_header_row(ws_new, ["Label", "Sublabel", "Source", "Description", "Created by"], FILL_HDR_NEW)
    for i, t in enumerate(new_labels, start=2):
        _nl(ws_new, i, 1, t.label)
        _nl(ws_new, i, 2, t.sublabel)
        _nl(ws_new, i, 3, t.source or "")
        _nl(ws_new, i, 4, t.description or "")
        _nl(ws_new, i, 5, creators.get(t.id, ""))
    _autofit(ws_new)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=labelbook_revised.xlsx"},
    )
