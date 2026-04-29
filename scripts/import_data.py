"""
Import user stories and label taxonomy from spreadsheets into the database.
Run once (or re-run with --reset to clear and reimport).

Usage:
    python -m scripts.import_data
    python -m scripts.import_data --reset
"""
import asyncio
import re
import sys
import random
from pathlib import Path

import openpyxl
from sqlalchemy import delete, select, text

sys.path.insert(0, str(Path(__file__).parent.parent))

from app.database import init_db, SessionLocal
from app.models import (
    Group, UserStory, StoryLabel, TaxonomyLabel, GroupAssignment
)

DATASET_PATH = Path("input_data/combined_output_VALIDATION_READY_EXPANDED_AI_LABELS_v007.xlsx")
LABELBOOK_PATH = Path("labelbook/Revised labelbook proposal for Oulu.xlsx")

GROUP_NAMES = ["Group A", "Group B", "Group C"]
# Each group reviews this fraction of all stories; overlap comes from the random draws
SUBSET_FRACTION = 0.6
RANDOM_SEED = 42


def parse_stories(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]

    def col(name: str) -> int:
        for i, h in enumerate(headers):
            if h and h.strip() == name:
                return i
        raise KeyError(f"Column not found: {name!r}")

    idx_story_id     = col("User Story ID")
    idx_workshop     = col("Workshop / Engagement Session")
    idx_submitted_by = col("Who submitted it")
    idx_stakeholder  = col("Stakeholder Group")
    idx_user_type    = col("User type")
    idx_task         = col("Task")
    idx_goal         = col("Goal")
    idx_notes        = col("Additional Notes (optional)")

    # Collect human label columns: (label_col, status_col, label_num)
    human_pairs: list[tuple[int, int | None, int]] = []
    ai_cols: list[tuple[int, int]] = []
    for i, h in enumerate(headers):
        if not h:
            continue
        m = re.match(r'^Human label (\d+)$', h.strip(), re.IGNORECASE)
        if m:
            num = int(m.group(1))
            status_name = f"Human label {num} Status"
            status_idx = next((j for j, sh in enumerate(headers)
                               if sh and sh.strip() == status_name), None)
            human_pairs.append((i, status_idx, num))
            continue
        m2 = re.match(r'^AI label (\d+)$', h.strip(), re.IGNORECASE)
        if m2:
            ai_cols.append((i, int(m2.group(1))))

    seen_ids = set()
    stories = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        story_id = row[idx_story_id]
        if not story_id or story_id in seen_ids:
            continue
        seen_ids.add(story_id)

        raw_labels = []
        for label_idx, status_idx, num in human_pairs:
            val = row[label_idx]
            if not val:
                continue
            status = row[status_idx] if status_idx is not None else None
            for text_val in str(val).split(";"):
                text_val = text_val.strip()
                if text_val:
                    raw_labels.append({
                        "source": "Human",
                        "label_text": text_val,
                        "label_index": num,
                        "label_status": status,
                    })
        for label_idx, num in ai_cols:
            val = row[label_idx]
            if not val:
                continue
            for text_val in str(val).split(";"):
                text_val = text_val.strip()
                if text_val:
                    raw_labels.append({
                        "source": "AI",
                        "label_text": text_val,
                        "label_index": num,
                        "label_status": None,
                    })

        stories.append({
            "story_id": story_id,
            "workshop": row[idx_workshop],
            "submitted_by": row[idx_submitted_by],
            "stakeholder_group": row[idx_stakeholder],
            "user_type": row[idx_user_type],
            "task": row[idx_task],
            "goal": row[idx_goal],
            "additional_notes": row[idx_notes],
            "labels": raw_labels,
        })

    return stories


def parse_taxonomy(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path)
    ws = wb.active

    taxonomy = []
    current_label = None

    # Main taxonomy starts at row 10 (row 10 is the sub-header)
    for row in ws.iter_rows(min_row=11, max_row=ws.max_row,
                             min_col=3, max_col=6, values_only=True):
        label_col, sublabel_col, source_col, desc_col = row

        if label_col:
            current_label = str(label_col).strip()
            # Top-level entry with no sublabel
            taxonomy.append({
                "label": current_label,
                "sublabel": None,
                "source": str(source_col).strip() if source_col else None,
                "description": str(desc_col).strip() if desc_col else None,
            })
        elif sublabel_col and current_label:
            taxonomy.append({
                "label": current_label,
                "sublabel": str(sublabel_col).strip(),
                "source": str(source_col).strip() if source_col else None,
                "description": str(desc_col).strip() if desc_col else None,
            })

    return taxonomy


def assign_subsets(story_ids: list[int], group_count: int,
                   fraction: float, seed: int) -> dict[int, list[int]]:
    """Give each group a random subset of story IDs with partial overlap."""
    rng = random.Random(seed)
    n = len(story_ids)
    subset_size = max(1, int(n * fraction))
    assignments: dict[int, list[int]] = {}
    for g in range(group_count):
        sample = rng.sample(story_ids, subset_size)
        assignments[g] = sample
    return assignments


async def run(reset: bool = False):
    await init_db()

    async with SessionLocal() as db:
        if reset:
            for table in [GroupAssignment, StoryLabel, UserStory,
                          TaxonomyLabel, Group]:
                await db.execute(delete(table))
            await db.commit()
            print("Database cleared.")

        # --- Groups ---
        existing_groups = (await db.execute(select(Group))).scalars().all()
        if not existing_groups:
            for name in GROUP_NAMES:
                db.add(Group(name=name))
            await db.commit()
            print(f"Created {len(GROUP_NAMES)} groups.")

        groups = (await db.execute(select(Group))).scalars().all()
        group_map = {g.name: g.id for g in groups}

        # --- Stories ---
        existing_stories = (await db.execute(select(UserStory))).scalars().all()
        if not existing_stories:
            stories = parse_stories(DATASET_PATH)
            for s in stories:
                story = UserStory(
                    story_id=s["story_id"],
                    workshop=s["workshop"],
                    submitted_by=s["submitted_by"],
                    stakeholder_group=s["stakeholder_group"],
                    user_type=s["user_type"],
                    task=s["task"],
                    goal=s["goal"],
                    additional_notes=s["additional_notes"],
                )
                db.add(story)
                await db.flush()
                for lbl in s["labels"]:
                    db.add(StoryLabel(
                        story_id=story.id,
                        source=lbl["source"],
                        label_text=lbl["label_text"],
                        label_index=lbl["label_index"],
                        label_status=lbl["label_status"],
                    ))
            await db.commit()
            print(f"Imported {len(stories)} user stories.")
        else:
            print(f"Stories already in DB ({len(existing_stories)}), skipping import.")

        # --- Taxonomy ---
        existing_tax = (await db.execute(select(TaxonomyLabel))).scalars().all()
        if not existing_tax:
            taxonomy = parse_taxonomy(LABELBOOK_PATH)
            for t in taxonomy:
                db.add(TaxonomyLabel(**t))
            await db.commit()
            print(f"Imported {len(taxonomy)} taxonomy entries.")
        else:
            print(f"Taxonomy already in DB ({len(existing_tax)}), skipping import.")

        # --- Group assignments ---
        existing_assignments = (await db.execute(select(GroupAssignment))).scalars().all()
        if not existing_assignments:
            all_stories = (await db.execute(select(UserStory))).scalars().all()
            story_ids = [s.id for s in all_stories]
            assignments = assign_subsets(story_ids, len(groups), SUBSET_FRACTION, RANDOM_SEED)

            for g_idx, group in enumerate(groups):
                for position, sid in enumerate(assignments[g_idx]):
                    db.add(GroupAssignment(
                        group_id=group.id,
                        story_id=sid,
                        position=position,
                    ))
            await db.commit()
            print("Group subsets assigned.")
        else:
            print("Group assignments already exist, skipping.")

    print("Done.")


if __name__ == "__main__":
    reset = "--reset" in sys.argv
    asyncio.run(run(reset=reset))
